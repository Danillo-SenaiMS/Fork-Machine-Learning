import pandas as pd
import matplotlib.pyplot as plt

# ======================
# 1. Base de dados fictícia
# ======================
data = {
    "Cliente": ["Cliente A", "Cliente B", "Cliente C", "Cliente D", "Cliente E"],
    "Faturamento": [500000, 200000, 150000, 80000, 30000],
    "Frequencia": [25, 10, 8, 5, 2],
    "Mix_Compras": [10, 7, 5, 3, 1],
    "Margem_Lucro": [0.25, 0.20, 0.18, 0.15, 0.10],
    "Ticket_Medio": [20000, 20000, 18750, 16000, 15000],
    "Tempo_Relacionamento": [5, 3, 2, 1, 1],
    "Risco_Inadimplencia": [1, 2, 3, 4, 5],
    "Custo_Atendimento": [2, 3, 4, 5, 5],
    "Engajamento": [9, 7, 5, 3, 2],
    "Sazonalidade": [8, 6, 5, 3, 2]
}
df = pd.DataFrame(data)

# ======================
# 2. Pesos definidos
# ======================
pesos = {
    "Faturamento": 0.15,
    "Frequencia": 0.15,
    "Mix_Compras": 0.10,
    "Margem_Lucro": 0.20,
    "Ticket_Medio": 0.08,
    "Tempo_Relacionamento": 0.10,
    "Risco_Inadimplencia": -0.02,
    "Custo_Atendimento": -0.04,
    "Engajamento": 0.10,
    "Sazonalidade": 0.06
}

# ======================
# 3. Normalização (0-10)
# ======================
for col in pesos.keys():
    if col in ["Risco_Inadimplencia", "Custo_Atendimento"]:
        df[col+"_Score"] = 10 - (df[col]-df[col].min())/(df[col].max()-df[col].min())*10
    else:
        df[col+"_Score"] = (df[col]-df[col].min())/(df[col].max()-df[col].min())*10

# ======================
# 4. Score final + Classificação
# ======================
df["Score_Final"] = sum(df[col+"_Score"]*peso for col, peso in pesos.items())

def classificar(score):
    if score >= 8: return "Classe A"
    elif score >= 6: return "Classe B"
    elif score >= 4: return "Classe C"
    else: return "Classe D"

df["Classe"] = df["Score_Final"].apply(classificar)

# ======================
# 5. Criar cenários simulados
# ======================
cenarios = df.copy()
cenarios["Score_Otimista"] = cenarios["Score_Final"] * 1.1
cenarios["Score_Pessimista"] = cenarios["Score_Final"] * 0.9

# ======================
# 6. Criar gráficos
# ======================
# Gráfico 1 - Distribuição por Classe (Pizza)
fig1, ax1 = plt.subplots()
df["Classe"].value_counts().plot.pie(autopct="%.1f%%", ax=ax1)
plt.title("Distribuição de Clientes por Classe")

# Gráfico 2 - TOP 10 clientes (Barras)
fig2, ax2 = plt.subplots()
top10 = df.nlargest(10, "Score_Final")[["Cliente","Score_Final"]]
ax2.bar(top10["Cliente"], top10["Score_Final"])
plt.title("TOP 10 Clientes por Score")
plt.xticks(rotation=45)

# Salvar gráficos como imagens temporárias
fig1.savefig("/mnt/data/distribuicao_clientes.png")
fig2.savefig("/mnt/data/top10_clientes.png")

plt.close(fig1)
plt.close(fig2)

# ======================
# 7. Exportar Excel com gráficos embutidos
# ======================
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image

wb = Workbook()

# Aba Base de Dados
ws1 = wb.active
ws1.title = "Base de Dados"
for r in dataframe_to_rows(pd.DataFrame(data), index=False, header=True):
    ws1.append(r)

# Aba Cálculo Score
ws2 = wb.create_sheet("Calculo Score")
for r in dataframe_to_rows(df, index=False, header=True):
    ws2.append(r)

# Aba Resumo
resumo = df.groupby("Classe").agg(Clientes=("Cliente","count"), Score_Medio=("Score_Final","mean")).reset_index()
ws3 = wb.create_sheet("Resumo")
for r in dataframe_to_rows(resumo, index=False, header=True):
    ws3.append(r)

# Aba Cenários
ws4 = wb.create_sheet("Cenarios")
for r in dataframe_to_rows(cenarios, index=False, header=True):
    ws4.append(r)

# Aba Gráficos
ws5 = wb.create_sheet("Graficos")
img1 = Image("/mnt/data/distribuicao_clientes.png")
img2 = Image("/mnt/data/top10_clientes.png")
ws5.add_image(img1, "B2")
ws5.add_image(img2, "B20")

# Salvar arquivo final
file_path = "/mnt/data/Segmentacao_Clientes_Completa.xlsx"
wb.save(file_path)

file_path
